from __future__ import annotations

import configparser
import json
import ssl
from base64 import b64encode
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable
from urllib import error, parse, request

from openpyxl import load_workbook


SCIM_GROUP_SCHEMA = "urn:ietf:params:scim:schemas:core:2.0:Group"
SCIM_USER_SCHEMA = "urn:ietf:params:scim:schemas:core:2.0:User"


@dataclass(frozen=True)
class AppConfig:
    tenant_url: str
    token_url: str
    client_id: str
    client_secret: str
    verify_ssl: bool = True
    timeout_seconds: int = 30
    scim_base_url: str = ""


@dataclass(frozen=True)
class TeamRow:
    team_id: str
    description: str


@dataclass(frozen=True)
class TeamRoleRow:
    team_id: str
    role_id: str


@dataclass(frozen=True)
class UserRow:
    username: str
    team_id: str


@dataclass(frozen=True)
class WorkbookData:
    teams: list[TeamRow]
    team_roles: list[TeamRoleRow]
    users: list[UserRow]


@dataclass(frozen=True)
class CreateTeamOp:
    team_id: str
    description: str


@dataclass(frozen=True)
class AssignRoleOp:
    team_id: str
    role_ids: list[str]


@dataclass(frozen=True)
class AssignUserOp:
    username: str
    team_ids: list[str]


@dataclass(frozen=True)
class ExecutionPlan:
    create_team_ops: list[CreateTeamOp]
    assign_role_ops: list[AssignRoleOp]
    assign_user_ops: list[AssignUserOp]


def workbook_record_counts(data: WorkbookData) -> dict[str, int]:
    return {
        "create_teams": len(data.teams),
        "assign_roles": len(data.team_roles),
        "users": len(data.users),
    }


def _clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_bool(value: str, default: bool = True) -> bool:
    text = _clean(value).lower()
    if text == "":
        return default
    return text in {"1", "true", "yes", "y"}


def load_config(path: Path) -> AppConfig:
    parser = configparser.ConfigParser()
    read_files = parser.read(path, encoding="utf-8")
    if not read_files:
        raise FileNotFoundError(f"Config file not found: {path}")
    if "SAC" not in parser:
        raise ValueError("config.ini must contain a [SAC] section.")

    sac = parser["SAC"]
    required = ["tenant_url", "token_url", "client_id", "client_secret"]
    missing = [field for field in required if _clean(sac.get(field, "")) == ""]
    if missing:
        raise ValueError(f"config.ini is missing required SAC fields: {', '.join(missing)}")

    tenant_url = _clean(sac["tenant_url"]).rstrip("/")
    scim_base_url = _clean(sac.get("scim_base_url")) or f"{tenant_url}/api/v1/scim"

    return AppConfig(
        tenant_url=tenant_url,
        token_url=_clean(sac["token_url"]),
        client_id=_clean(sac["client_id"]),
        client_secret=_clean(sac["client_secret"]),
        verify_ssl=_to_bool(sac.get("verify_ssl", "true"), default=True),
        timeout_seconds=int(_clean(sac.get("timeout_seconds")) or "30"),
        scim_base_url=scim_base_url.rstrip("/"),
    )


def save_config(config: AppConfig, path: Path) -> None:
    parser = configparser.ConfigParser()
    parser["SAC"] = {
        "tenant_url": config.tenant_url,
        "token_url": config.token_url,
        "client_id": config.client_id,
        "client_secret": config.client_secret,
    }
    with path.open("w", encoding="utf-8") as handle:
        parser.write(handle)


def _sheet_rows(path: Path, sheet_name: str):
    workbook = load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing worksheet '{sheet_name}' in {path.name}.")
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Worksheet '{sheet_name}' is empty.")
    header = [_clean(cell) for cell in rows[0]]
    for raw_row in rows[1:]:
        if all(_clean(cell) == "" for cell in raw_row):
            continue
        yield {header[index]: raw_row[index] if index < len(raw_row) else "" for index in range(len(header))}


def _assert_columns(path: Path, sheet_name: str, required_columns: Iterable[str]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing worksheet '{sheet_name}' in {path.name}.")
    sheet = workbook[sheet_name]
    header_row = next(sheet.iter_rows(values_only=True), None)
    if header_row is None:
        raise ValueError(f"Worksheet '{sheet_name}' is empty.")
    headers = {_clean(cell) for cell in header_row}
    missing = [column for column in required_columns if column not in headers]
    if missing:
        raise ValueError(
            f"Worksheet '{sheet_name}' is missing required columns: {', '.join(missing)}"
        )


def load_workbook_data(path: Path) -> WorkbookData:
    _assert_columns(path, "Create_Teams", ["Team ID", "Team Description"])
    _assert_columns(path, "Assign_Roles", ["TeamID", "RoleID"])
    _assert_columns(path, "Users", ["UserName", "TeamID"])

    teams = [
        TeamRow(team_id=_clean(row["Team ID"]), description=_clean(row.get("Team Description")))
        for row in _sheet_rows(path, "Create_Teams")
        if _clean(row.get("Team ID"))
    ]
    team_roles = [
        TeamRoleRow(team_id=_clean(row["TeamID"]), role_id=_clean(row["RoleID"]))
        for row in _sheet_rows(path, "Assign_Roles")
        if _clean(row.get("TeamID")) and _clean(row.get("RoleID"))
    ]
    users = [
        UserRow(username=_clean(row["UserName"]), team_id=_clean(row["TeamID"]))
        for row in _sheet_rows(path, "Users")
        if _clean(row.get("UserName")) and _clean(row.get("TeamID"))
    ]

    return WorkbookData(teams=teams, team_roles=team_roles, users=users)


def load_workbook_data_from_bytes(payload: bytes) -> WorkbookData:
    with BytesIO(payload) as stream:
        workbook = load_workbook(stream, data_only=True)

    def sheet_rows(sheet_name: str):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet '{sheet_name}' in uploaded workbook.")
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"Worksheet '{sheet_name}' is empty.")
        header = [_clean(cell) for cell in rows[0]]
        for raw_row in rows[1:]:
            if all(_clean(cell) == "" for cell in raw_row):
                continue
            yield {
                header[index]: raw_row[index] if index < len(raw_row) else ""
                for index in range(len(header))
            }

    def assert_columns(sheet_name: str, required_columns: Iterable[str]) -> None:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet '{sheet_name}' in uploaded workbook.")
        sheet = workbook[sheet_name]
        header_row = next(sheet.iter_rows(values_only=True), None)
        if header_row is None:
            raise ValueError(f"Worksheet '{sheet_name}' is empty.")
        headers = {_clean(cell) for cell in header_row}
        missing = [column for column in required_columns if column not in headers]
        if missing:
            raise ValueError(
                f"Worksheet '{sheet_name}' is missing required columns: {', '.join(missing)}"
            )

    assert_columns("Create_Teams", ["Team ID", "Team Description"])
    assert_columns("Assign_Roles", ["TeamID", "RoleID"])
    assert_columns("Users", ["UserName", "TeamID"])

    teams = [
        TeamRow(team_id=_clean(row["Team ID"]), description=_clean(row.get("Team Description")))
        for row in sheet_rows("Create_Teams")
        if _clean(row.get("Team ID"))
    ]
    team_roles = [
        TeamRoleRow(team_id=_clean(row["TeamID"]), role_id=_clean(row["RoleID"]))
        for row in sheet_rows("Assign_Roles")
        if _clean(row.get("TeamID")) and _clean(row.get("RoleID"))
    ]
    users = [
        UserRow(username=_clean(row["UserName"]), team_id=_clean(row["TeamID"]))
        for row in sheet_rows("Users")
        if _clean(row.get("UserName")) and _clean(row.get("TeamID"))
    ]

    return WorkbookData(teams=teams, team_roles=team_roles, users=users)


def validate_workbook_data(data: WorkbookData) -> list[str]:
    errors: list[str] = []

    seen_teams: set[str] = set()
    for team in data.teams:
        if team.team_id in seen_teams:
            errors.append(f"Create_Teams contains duplicate Team ID '{team.team_id}'.")
        seen_teams.add(team.team_id)

    for role_row in data.team_roles:
        if role_row.team_id == "":
            errors.append("Assign_Roles contains an empty TeamID.")
        if role_row.role_id == "":
            errors.append(f"Assign_Roles contains an empty RoleID for TeamID '{role_row.team_id}'.")

    for user_row in data.users:
        if user_row.username == "":
            errors.append("Users contains an empty UserName.")
        if user_row.team_id == "":
            errors.append(f"Users contains an empty TeamID for user '{user_row.username}'.")

    return errors


def normalize_role_id(role_id: str) -> str:
    return role_id


def build_execution_plan(data: WorkbookData) -> ExecutionPlan:
    create_team_ops = [
        CreateTeamOp(team_id=team.team_id, description=team.description) for team in data.teams
    ]

    grouped_roles: dict[str, list[str]] = {}
    for row in data.team_roles:
        grouped_roles.setdefault(row.team_id, [])
        grouped_roles[row.team_id].append(normalize_role_id(row.role_id))
    assign_role_ops = [
        AssignRoleOp(team_id=team_id, role_ids=role_ids) for team_id, role_ids in grouped_roles.items()
    ]

    grouped_users: dict[str, list[str]] = {}
    for row in data.users:
        grouped_users.setdefault(row.username, [])
        grouped_users[row.username].append(row.team_id)
    assign_user_ops = [
        AssignUserOp(username=username, team_ids=team_ids)
        for username, team_ids in grouped_users.items()
    ]

    return ExecutionPlan(
        create_team_ops=create_team_ops,
        assign_role_ops=assign_role_ops,
        assign_user_ops=assign_user_ops,
    )


def dry_run_report(plan: ExecutionPlan) -> str:
    lines = [
        "DRY RUN: no SAC changes were sent.",
        "",
        f"CREATE TEAMS ({len(plan.create_team_ops)})",
    ]
    for op in plan.create_team_ops:
        lines.append(f"- {op.team_id}: {op.description}")

    lines.append("")
    lines.append(f"ASSIGN ROLES ({len(plan.assign_role_ops)})")
    for op in plan.assign_role_ops:
        lines.append(f"- {op.team_id}: {', '.join(op.role_ids)}")

    lines.append("")
    lines.append(f"ASSIGN USERS ({len(plan.assign_user_ops)})")
    for op in plan.assign_user_ops:
        lines.append(f"- {op.username}: {', '.join(op.team_ids)}")
    return "\n".join(lines)


class SacScimClient:
    """Our implementation for SAC SCIM provisioning.

    This wraps the existing repo approach, but keeps credentials out of code and
    centralizes the HTTP assumptions in one place.
    """

    def __init__(self, config: AppConfig):
        self.config = config
        self._token: str | None = None
        self._csrf_cache: dict[str, tuple[str | None, str | None]] = {}

    def _ssl_context(self):
        if self.config.verify_ssl:
            return None
        return ssl._create_unverified_context()

    def _http(self, req: request.Request) -> tuple[int, dict[str, str], str]:
        try:
            with request.urlopen(
                req,
                timeout=self.config.timeout_seconds,
                context=self._ssl_context(),
            ) as response:
                body = response.read().decode("utf-8")
                return response.status, dict(response.headers.items()), body
        except error.HTTPError as exc:
            body = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"HTTP {exc.code} for {req.full_url}: {body}") from exc

    def get_access_token(self) -> str:
        if self._token:
            return self._token

        token_url = self.config.token_url
        separator = "&" if "?" in token_url else "?"
        token_url = f"{token_url}{separator}grant_type=client_credentials"
        req = request.Request(token_url, method="POST")
        basic = b64encode(
            f"{self.config.client_id}:{self.config.client_secret}".encode("utf-8")
        ).decode("ascii")
        req.add_header("Authorization", f"Basic {basic}")
        status, _, body = self._http(req)
        if status != 200:
            raise RuntimeError(f"Token request failed with status {status}.")
        parsed = json.loads(body)
        token = parsed.get("access_token")
        if not token:
            raise RuntimeError("Token response did not include access_token.")
        self._token = token
        return token

    def _base_headers(self) -> dict[str, str]:
        return {
            "Authorization": f"Bearer {self.get_access_token()}",
            "x-sap-sac-custom-auth": "true",
            "Accept": "application/scim+json",
        }

    def fetch_csrf(self, resource: str) -> tuple[str | None, str | None]:
        if resource in self._csrf_cache:
            return self._csrf_cache[resource]
        req = request.Request(f"{self.config.scim_base_url}/{resource}", method="GET")
        for key, value in self._base_headers().items():
            req.add_header(key, value)
        req.add_header("x-csrf-token", "fetch")
        status, headers, _ = self._http(req)
        if status != 200:
            raise RuntimeError(f"Unable to fetch CSRF token for {resource}.")
        csrf = headers.get("x-csrf-token")
        cookie = headers.get("set-cookie")
        self._csrf_cache[resource] = (csrf, cookie)
        return csrf, cookie

    def _json_request(self, method: str, url: str, payload: dict, csrf_resource: str) -> dict:
        body = json.dumps(payload).encode("utf-8")
        req = request.Request(url, data=body, method=method)
        for key, value in self._base_headers().items():
            req.add_header(key, value)
        req.add_header("Content-Type", "application/scim+json")
        csrf, cookie = self.fetch_csrf(csrf_resource)
        if csrf:
            req.add_header("x-csrf-token", csrf)
        if cookie:
            req.add_header("Cookie", cookie)
        _, _, text = self._http(req)
        return json.loads(text) if text else {}

    def get_team(self, team_id: str) -> dict | None:
        req = request.Request(f"{self.config.scim_base_url}/Groups/{parse.quote(team_id)}", method="GET")
        for key, value in self._base_headers().items():
            req.add_header(key, value)
        try:
            _, _, body = self._http(req)
        except RuntimeError as exc:
            if "HTTP 404" in str(exc):
                return None
            raise
        return json.loads(body) if body else None

    def get_user(self, username: str) -> dict | None:
        req = request.Request(f"{self.config.scim_base_url}/Users/{parse.quote(username)}", method="GET")
        for key, value in self._base_headers().items():
            req.add_header(key, value)
        try:
            _, _, body = self._http(req)
        except RuntimeError as exc:
            if "HTTP 404" in str(exc):
                return None
            raise
        return json.loads(body) if body else None

    def create_team(self, op: CreateTeamOp) -> None:
        existing = self.get_team(op.team_id)
        if existing:
            return
        payload = {
            "schemas": [SCIM_GROUP_SCHEMA],
            "id": op.team_id,
            "displayName": op.description or op.team_id,
        }
        self._json_request("POST", f"{self.config.scim_base_url}/Groups", payload, "Groups")

    def assign_roles(self, op: AssignRoleOp) -> None:
        team = self.get_team(op.team_id)
        if not team:
            raise RuntimeError(f"Team '{op.team_id}' does not exist.")
        team["roles"] = op.role_ids
        self._json_request(
            "PUT",
            f"{self.config.scim_base_url}/Groups/{parse.quote(op.team_id)}",
            team,
            "Groups",
        )

    def assign_user_teams(self, op: AssignUserOp) -> None:
        user = self.get_user(op.username)
        if not user:
            raise RuntimeError(
                f"User '{op.username}' does not exist in SAC. "
                "This workbook only provides UserName and TeamID, so the tool can assign existing users only."
            )
        existing_team_ids = {
            entry.get("value") for entry in user.get("groups", []) if entry.get("value")
        }
        user["groups"] = [{"value": team_id} for team_id in sorted(existing_team_ids.union(op.team_ids))]
        self._json_request(
            "PUT",
            f"{self.config.scim_base_url}/Users/{parse.quote(op.username)}",
            user,
            "Users",
        )


def run_create_teams(client: SacScimClient, plan: ExecutionPlan) -> list[str]:
    lines = []
    for op in plan.create_team_ops:
        client.create_team(op)
        lines.append(f"Team ready: {op.team_id}")
    return lines


def run_assign_roles(client: SacScimClient, plan: ExecutionPlan) -> list[str]:
    lines = []
    for op in plan.assign_role_ops:
        client.assign_roles(op)
        lines.append(f"Roles assigned: {op.team_id} -> {', '.join(op.role_ids)}")
    return lines


def run_assign_users(client: SacScimClient, plan: ExecutionPlan) -> list[str]:
    lines = []
    for op in plan.assign_user_ops:
        client.assign_user_teams(op)
        lines.append(f"User assigned: {op.username} -> {', '.join(op.team_ids)}")
    return lines
